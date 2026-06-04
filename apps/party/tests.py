import shutil
from io import BytesIO
from pathlib import Path

from django.core.management import call_command
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.conf import settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.party.models import (
    PartyApprovalTask,
    PartyMaterialSubmission,
    PartyMaterialType,
    PartyMemberStatus,
    PartyReminder,
    PartyStageDefinition,
    PartyStageInstance,
    PartyStageMaterialRequirement,
)
from apps.party.services import approval, reminder, workflow
from apps.users.models import AuditLog, User


class PartyWorkflowTestCase(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._temp_media = Path(settings.BASE_DIR) / ".tmp_party_test_media"
        cls._temp_media.mkdir(parents=True, exist_ok=True)
        cls._override = override_settings(MEDIA_ROOT=cls._temp_media)
        cls._override.enable()

    @classmethod
    def tearDownClass(cls):
        cls._override.disable()
        shutil.rmtree(cls._temp_media, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.admin_user = User.objects.create_user(
            username="admin_user",
            password="party1234",
            role=User.ROLE_ADMIN,
            real_name="管理老师",
        )
        self.student_user = User.objects.create_user(
            username="student_user",
            password="party1234",
            role=User.ROLE_STUDENT,
            real_name="学生甲",
            student_id="20260011",
        )
        self.material_type = PartyMaterialType.objects.create(
            code="application_letter",
            name="入党申请书",
            allowed_extensions="pdf,doc,docx",
            max_size_mb=30,
        )
        self.stage_one = PartyStageDefinition.objects.create(
            code="party_applicant",
            name="入党申请人",
            track_type=PartyMemberStatus.TrackType.PARTY,
            stage=PartyMemberStatus.Stage.APPLICANT,
            order=1,
            target_role_for_approval=User.ROLE_ADMIN,
        )
        self.stage_two = PartyStageDefinition.objects.create(
            code="party_activist",
            name="积极分子",
            track_type=PartyMemberStatus.TrackType.PARTY,
            stage=PartyMemberStatus.Stage.ACTIVIST,
            order=2,
            min_duration_days=90,
            target_role_for_approval=User.ROLE_ADMIN,
        )
        PartyStageMaterialRequirement.objects.create(
            stage_definition=self.stage_one,
            material_type=self.material_type,
            is_required=True,
            order=1,
        )
        self.member = PartyMemberStatus.objects.create(
            user=self.student_user,
            track_type=PartyMemberStatus.TrackType.PARTY,
            current_stage=PartyMemberStatus.Stage.APPLICANT,
            created_by=self.admin_user,
        )
        self.current_stage = workflow.initialize_member_workflow(
            self.member,
            operator=self.admin_user,
        )

    def _submit_required_material(self):
        upload = SimpleUploadedFile(
            "application.pdf",
            b"test content",
            content_type="application/pdf",
        )
        return PartyMaterialSubmission.objects.create(
            stage_instance=self.current_stage,
            material_type=self.material_type,
            file=upload,
            version=workflow.get_next_material_version(self.current_stage, self.material_type),
            submitted_by=self.student_user,
        )

    def _create_pending_task_for_member(self, username, real_name, student_id):
        user = User.objects.create_user(
            username=username,
            password="party1234",
            role=User.ROLE_STUDENT,
            real_name=real_name,
            student_id=student_id,
        )
        member = PartyMemberStatus.objects.create(
            user=user,
            track_type=PartyMemberStatus.TrackType.PARTY,
            current_stage=PartyMemberStatus.Stage.APPLICANT,
            created_by=self.admin_user,
        )
        stage_instance = workflow.initialize_member_workflow(member, operator=self.admin_user)
        PartyMaterialSubmission.objects.create(
            stage_instance=stage_instance,
            material_type=self.material_type,
            file=SimpleUploadedFile(
                f"{username}.pdf",
                b"test content",
                content_type="application/pdf",
            ),
            version=workflow.get_next_material_version(stage_instance, self.material_type),
            submitted_by=user,
        )
        task = workflow.submit_stage_for_approval(
            stage_instance,
            user,
            comment=f"{real_name} 提交审批",
        )
        return user, member, stage_instance, task

    def test_submit_approval_requires_all_required_materials(self):
        with self.assertRaisesMessage(ValidationError, "当前节点仍有必需材料未提交。"):
            workflow.submit_stage_for_approval(self.current_stage, self.student_user)

    def test_party_index_redirects_anonymous_entry_links_to_login(self):
        response = self.client.get("/party/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "/users/login/?next=/party/student/")
        self.assertContains(response, "/users/login/?next=/party/admin/")

    def test_login_redirects_to_requested_party_page(self):
        response = self.client.get("/users/login/?next=/party/student/timeline/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'name="next" value="/party/student/timeline/"')

        response = self.client.post(
            "/users/login/?next=/party/student/timeline/",
            {
                "username": "student_user",
                "password": "party1234",
                "next": "/party/student/timeline/",
            },
        )

        self.assertRedirects(
            response,
            "/party/student/timeline/",
            fetch_redirect_response=False,
        )

    def test_submit_approval_creates_pending_task(self):
        self._submit_required_material()

        task = workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="材料已上传完成",
        )

        self.current_stage.refresh_from_db()
        self.assertEqual(task.status, PartyApprovalTask.Status.PENDING)
        self.assertEqual(task.assignee_role, User.ROLE_ADMIN)
        self.assertEqual(self.current_stage.status, PartyStageInstance.Status.SUBMITTED)
        self.assertEqual(task.actions.count(), 1)
        self.assertEqual(task.actions.first().action, "submit")

    def test_approve_task_advances_workflow_to_next_stage(self):
        self._submit_required_material()
        task = workflow.submit_stage_for_approval(self.current_stage, self.student_user)

        next_stage = approval.approve_task(task, self.admin_user, comment="通过")

        task.refresh_from_db()
        self.current_stage.refresh_from_db()
        self.member.refresh_from_db()
        self.assertEqual(task.status, PartyApprovalTask.Status.APPROVED)
        self.assertEqual(self.current_stage.status, PartyStageInstance.Status.APPROVED)
        self.assertIsNotNone(next_stage)
        self.assertEqual(next_stage.stage_definition, self.stage_two)
        self.assertEqual(next_stage.status, PartyStageInstance.Status.OPEN)
        self.assertEqual(self.member.current_stage, PartyMemberStatus.Stage.ACTIVIST)

    def test_reject_task_returns_stage_to_rejected_state(self):
        self._submit_required_material()
        task = workflow.submit_stage_for_approval(self.current_stage, self.student_user)

        stage_instance = approval.reject_task(
            task,
            self.admin_user,
            comment="材料内容不完整",
        )

        task.refresh_from_db()
        stage_instance.refresh_from_db()
        self.assertEqual(task.status, PartyApprovalTask.Status.REJECTED)
        self.assertEqual(stage_instance.status, PartyStageInstance.Status.REJECTED)
        self.assertEqual(stage_instance.remark, "材料内容不完整")

    def test_student_history_page_renders_own_history(self):
        self._submit_required_material()
        task = workflow.submit_stage_for_approval(self.current_stage, self.student_user)
        approval.reject_task(task, self.admin_user, comment="请补交签名页")

        self.client.force_login(self.student_user)
        response = self.client.get("/party/student/history/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "入党申请人")
        self.assertContains(response, "入党申请书")
        self.assertContains(response, "请补交签名页")

    def test_student_history_page_highlights_withdrawn_chain(self):
        self._submit_required_material()
        task = workflow.submit_stage_for_approval(self.current_stage, self.student_user)
        approval.approve_task(task, self.admin_user, comment="审核通过")
        approval.withdraw_task(task, self.admin_user, comment="发现误审，重新审核")

        self.client.force_login(self.student_user)
        response = self.client.get("/party/student/history/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "原审批结果已撤回")
        self.assertContains(response, "撤回后状态：该节点已重新打开审核")
        self.assertContains(response, "发现误审，重新审核")

    def test_student_post_actions_create_request_and_business_audit_logs(self):
        self.client.force_login(self.student_user)

        upload = SimpleUploadedFile(
            "application.pdf",
            b"test content",
            content_type="application/pdf",
        )
        response = self.client.post(
            "/party/student/materials/",
            {
                "action": "upload_material",
                "material_type": self.material_type.pk,
                "file": upload,
            },
        )
        self.assertEqual(response.status_code, 302)

        self.current_stage.refresh_from_db()
        upload_logs = AuditLog.objects.filter(user=self.student_user)
        self.assertTrue(upload_logs.filter(action="POST /party/student/materials/").exists())
        self.assertTrue(upload_logs.filter(action="upload party material").exists())

        response = self.client.post(
            "/party/student/materials/",
            {
                "action": "submit_approval",
                "comment": "学生提交审批",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(
            AuditLog.objects.filter(
                user=self.student_user,
                action="submit party approval",
            ).exists()
        )

    def test_admin_approval_post_creates_request_and_business_audit_logs(self):
        self._submit_required_material()
        task = workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="学生提交审批",
        )

        self.client.force_login(self.admin_user)
        response = self.client.post(
            "/party/admin/approvals/",
            {
                "task_id": task.pk,
                "action": "approve",
                "comment": "审核通过",
            },
        )
        self.assertEqual(response.status_code, 302)

        admin_logs = AuditLog.objects.filter(user=self.admin_user)
        self.assertTrue(admin_logs.filter(action="POST /party/admin/approvals/").exists())
        self.assertTrue(admin_logs.filter(action="approve party task").exists())
        self.assertTrue(admin_logs.filter(action="advance party stage").exists())

    def test_admin_approval_page_can_batch_approve_tasks(self):
        self._submit_required_material()
        first_task = workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="学生提交审批",
        )
        _, second_member, second_stage_instance, second_task = self._create_pending_task_for_member(
            "batch_student",
            "批量审批学生",
            "20260013",
        )

        self.client.force_login(self.admin_user)
        response = self.client.post(
            "/party/admin/approvals/?q=学生",
            {
                "operation": "batch_process",
                "task_ids": [str(first_task.pk), str(second_task.pk)],
                "action": "approve",
                "comment": "批量审核通过",
                "redirect_to": "/party/admin/approvals/?q=%E5%AD%A6%E7%94%9F",
            },
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response["Location"], "/party/admin/approvals/?q=%E5%AD%A6%E7%94%9F")

        first_task.refresh_from_db()
        second_task.refresh_from_db()
        self.current_stage.refresh_from_db()
        second_stage_instance.refresh_from_db()
        self.member.refresh_from_db()
        second_member.refresh_from_db()

        self.assertEqual(first_task.status, PartyApprovalTask.Status.APPROVED)
        self.assertEqual(second_task.status, PartyApprovalTask.Status.APPROVED)
        self.assertEqual(self.current_stage.status, PartyStageInstance.Status.APPROVED)
        self.assertEqual(second_stage_instance.status, PartyStageInstance.Status.APPROVED)
        self.assertEqual(self.member.current_stage, PartyMemberStatus.Stage.ACTIVIST)
        self.assertEqual(second_member.current_stage, PartyMemberStatus.Stage.ACTIVIST)
        self.assertTrue(
            AuditLog.objects.filter(
                user=self.admin_user,
                action="batch approve party tasks",
            ).exists()
        )

    def test_admin_approval_page_batch_reject_requires_comment(self):
        self._submit_required_material()
        task = workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="学生提交审批",
        )

        self.client.force_login(self.admin_user)
        response = self.client.post(
            "/party/admin/approvals/",
            {
                "operation": "batch_process",
                "task_ids": [str(task.pk)],
                "action": "reject",
                "comment": "",
                "redirect_to": "/party/admin/approvals/",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "批量驳回时必须填写原因。")
        task.refresh_from_db()
        self.assertEqual(task.status, PartyApprovalTask.Status.PENDING)

    def test_withdraw_approved_task_reopens_pending_task(self):
        self._submit_required_material()
        task = workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="学生提交审批",
        )
        approval.approve_task(task, self.admin_user, comment="审核通过")

        task.refresh_from_db()
        reopened_task = approval.withdraw_task(
            task,
            self.admin_user,
            comment="撤回重审",
        )

        self.current_stage.refresh_from_db()
        self.member.refresh_from_db()
        self.assertEqual(task.status, PartyApprovalTask.Status.WITHDRAWN)
        self.assertEqual(reopened_task.status, PartyApprovalTask.Status.PENDING)
        self.assertEqual(self.current_stage.status, PartyStageInstance.Status.SUBMITTED)
        self.assertEqual(self.member.current_stage, PartyMemberStatus.Stage.APPLICANT)
        self.assertTrue(
            AuditLog.objects.filter(
                user=self.admin_user,
                action="withdraw party task",
            ).exists()
        )

    def test_withdraw_requires_comment(self):
        self._submit_required_material()
        task = workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="学生提交审批",
        )
        approval.approve_task(task, self.admin_user, comment="审核通过")
        task.refresh_from_db()

        with self.assertRaisesMessage(ValidationError, "撤回时必须填写说明。"):
            approval.withdraw_task(task, self.admin_user, comment="")

    def test_withdraw_only_allows_original_approver_or_leader(self):
        other_admin = User.objects.create_user(
            username="other_admin",
            password="party1234",
            role=User.ROLE_ADMIN,
            real_name="其他管理员",
        )
        leader_user = User.objects.create_user(
            username="leader_user",
            password="party1234",
            role=User.ROLE_LEADER,
            real_name="学院领导",
        )
        self._submit_required_material()
        task = workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="学生提交审批",
        )
        approval.approve_task(task, self.admin_user, comment="审核通过")
        task.refresh_from_db()

        with self.assertRaisesMessage(ValidationError, "仅原审批人或学院领导可以撤回该审批结果。"):
            approval.withdraw_task(task, other_admin, comment="越权撤回")

        reopened_task = approval.withdraw_task(task, leader_user, comment="领导要求重审")
        self.assertEqual(reopened_task.status, PartyApprovalTask.Status.PENDING)

    def test_withdraw_rejected_task_is_blocked_after_new_material_submission(self):
        self._submit_required_material()
        task = workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="学生提交审批",
        )
        approval.reject_task(task, self.admin_user, comment="请补材料")
        task.refresh_from_db()

        new_submission = PartyMaterialSubmission.objects.create(
            stage_instance=self.current_stage,
            material_type=self.material_type,
            file=SimpleUploadedFile(
                "reupload.pdf",
                b"new content",
                content_type="application/pdf",
            ),
            version=workflow.get_next_material_version(self.current_stage, self.material_type),
            submitted_by=self.student_user,
        )
        PartyMaterialSubmission.objects.filter(pk=new_submission.pk).update(
            submitted_at=task.processed_at + timezone.timedelta(minutes=1)
        )

        with self.assertRaisesMessage(ValidationError, "学生在驳回后已经补交新材料，当前结果不能撤回。"):
            approval.withdraw_task(task, self.admin_user, comment="想撤回驳回")

    def test_admin_approval_page_shows_task_chain_after_withdraw(self):
        self._submit_required_material()
        task = workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="学生提交审批",
        )
        approval.approve_task(task, self.admin_user, comment="审核通过")
        approval.withdraw_task(task, self.admin_user, comment="撤回重审")

        self.client.force_login(self.admin_user)
        response = self.client.get(f"/party/admin/approvals/?task={task.id}")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "当前节点审批链路")
        self.assertContains(response, "已撤回")
        self.assertContains(response, "已重新打开审核")

    def test_admin_member_detail_page_renders_full_member_context(self):
        self._submit_required_material()
        task = workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="学生提交审批",
        )
        approval.reject_task(task, self.admin_user, comment="请补交签名页")

        self.client.force_login(self.admin_user)
        response = self.client.get(f"/party/admin/members/{self.member.id}/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "成员详情")
        self.assertContains(response, "学生甲")
        self.assertContains(response, "流程时间轴")
        self.assertContains(response, "节点材料与审批链路")
        self.assertContains(response, "请补交签名页")

    def test_admin_member_detail_page_keeps_return_to_context(self):
        self.client.force_login(self.admin_user)

        return_to = "/party/admin/members/?q=student_user&stage=applicant"
        response = self.client.get(
            f"/party/admin/members/{self.member.id}/?return_to=%2Fparty%2Fadmin%2Fmembers%2F%3Fq%3Dstudent_user%26stage%3Dapplicant"
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            'href="/party/admin/members/?q=student_user&amp;stage=applicant"',
        )
        self.assertContains(response, "breadcrumb-list")

    def test_admin_members_page_supports_filters(self):
        other_student = User.objects.create_user(
            username="other_student",
            password="party1234",
            role=User.ROLE_STUDENT,
            real_name="学生乙",
            student_id="20260012",
            major="软件工程",
        )
        PartyMemberStatus.objects.create(
            user=other_student,
            track_type=PartyMemberStatus.TrackType.PARTY,
            current_stage=PartyMemberStatus.Stage.CANDIDATE,
            created_by=self.admin_user,
        )

        self.client.force_login(self.admin_user)
        response = self.client.get("/party/admin/members/?q=学生甲&stage=applicant&task_status=none")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "学生甲")
        self.assertNotContains(response, "学生乙")

    def test_admin_members_page_supports_pagination_and_sort(self):
        for index in range(12):
            user = User.objects.create_user(
                username=f"bulk_student_{index}",
                password="party1234",
                role=User.ROLE_STUDENT,
                real_name=f"批量学生{index:02d}",
                student_id=f"20261{index:03d}",
                grade="2026",
            )
            PartyMemberStatus.objects.create(
                user=user,
                track_type=PartyMemberStatus.TrackType.PARTY,
                current_stage=PartyMemberStatus.Stage.APPLICANT,
                created_by=self.admin_user,
            )

        self.client.force_login(self.admin_user)
        response = self.client.get("/party/admin/members/?sort=name_desc&page=2")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "第 2 / 2 页")
        self.assertContains(response, "批量学生01")

    def test_admin_member_detail_links_to_filtered_approvals(self):
        self._submit_required_material()
        workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="学生提交审批",
        )

        self.client.force_login(self.admin_user)
        response = self.client.get(f"/party/admin/approvals/?member={self.member.id}")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "当前仅展示指定学生的相关审批任务")
        self.assertContains(response, "学生甲")

    def test_admin_approvals_page_supports_filters_and_pagination(self):
        self._submit_required_material()
        workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="学生甲审批",
        )
        for index in range(9):
            user = User.objects.create_user(
                username=f"approval_student_{index}",
                password="party1234",
                role=User.ROLE_STUDENT,
                real_name=f"审批学生{index}",
                student_id=f"30360{index}",
            )
            member = PartyMemberStatus.objects.create(
                user=user,
                track_type=PartyMemberStatus.TrackType.PARTY,
                current_stage=PartyMemberStatus.Stage.APPLICANT,
                created_by=self.admin_user,
            )
            stage_instance = workflow.initialize_member_workflow(member, operator=self.admin_user)
            PartyMaterialSubmission.objects.create(
                stage_instance=stage_instance,
                material_type=self.material_type,
                file=SimpleUploadedFile(
                    f"approval_{index}.pdf",
                    b"content",
                    content_type="application/pdf",
                ),
                version=1,
                submitted_by=user,
            )
            workflow.submit_stage_for_approval(stage_instance, user, comment=f"审批任务{index}")

        self.client.force_login(self.admin_user)
        response = self.client.get(
            "/party/admin/approvals/?q=审批学生&stage=applicant&sort=student_desc&pending_page=2"
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "第 2 / 2 页")
        self.assertContains(response, "审批学生0")

    def test_admin_approvals_page_renders_dense_summary_and_detail_fields(self):
        self._submit_required_material()
        workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="学生提交的补充说明",
        )

        self.client.force_login(self.admin_user)
        response = self.client.get("/party/admin/approvals/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "流程类型")
        self.assertContains(response, "材料完成度")
        self.assertContains(response, "账号")
        self.assertContains(response, "student_user")
        self.assertContains(response, "学生提交意见")
        self.assertContains(response, "学生提交的补充说明")
        self.assertContains(response, "材料：1/1")

    def test_admin_approvals_page_renders_decision_summary_and_member_context(self):
        self._submit_required_material()
        first_task = workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="第一次提交审批",
        )
        approval.reject_task(first_task, self.admin_user, comment="请补交签名页")
        second_upload = SimpleUploadedFile(
            "application_v2.pdf",
            b"new content",
            content_type="application/pdf",
        )
        PartyMaterialSubmission.objects.create(
            stage_instance=self.current_stage,
            material_type=self.material_type,
            file=second_upload,
            version=workflow.get_next_material_version(self.current_stage, self.material_type),
            submitted_by=self.student_user,
        )
        workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="第二次提交审批",
        )

        self.client.force_login(self.admin_user)
        response = self.client.get("/party/admin/approvals/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "成员上下文")
        self.assertContains(response, "决策摘要")
        self.assertContains(response, "当前节点历史任务数")
        self.assertContains(response, "最近一次驳回原因")
        self.assertContains(response, "请补交签名页")
        self.assertContains(response, "流程进度")
        self.assertContains(response, "年级 / 专业")
        self.assertContains(response, "成员待办")

    def test_admin_rules_page_renders_current_configuration(self):
        self.client.force_login(self.admin_user)

        response = self.client.get("/party/admin/rules/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "流程规则配置")
        self.assertContains(response, "节点定义")
        self.assertContains(response, "材料类型")
        self.assertContains(response, "入党申请人")
        self.assertContains(response, "入党申请书")

    def test_admin_rules_page_can_create_stage_definition(self):
        self.client.force_login(self.admin_user)

        response = self.client.post(
            "/party/admin/rules/",
            {
                "operation": "save_stage",
                "code": "party_candidate",
                "name": "发展对象",
                "track_type": PartyMemberStatus.TrackType.PARTY,
                "stage": PartyMemberStatus.Stage.CANDIDATE,
                "order": 3,
                "min_duration_days": 180,
                "target_role_for_approval": User.ROLE_ADMIN,
                "requires_training": "on",
                "is_active": "on",
                "description": "发展对象阶段规则",
            },
        )

        self.assertEqual(response.status_code, 302)
        created_stage = PartyStageDefinition.objects.get(code="party_candidate")
        self.assertEqual(created_stage.name, "发展对象")
        self.assertTrue(created_stage.requires_training)
        self.assertTrue(
            AuditLog.objects.filter(
                user=self.admin_user,
                action="save party stage definition",
                target_id=str(created_stage.id),
            ).exists()
        )

    def test_admin_rules_page_can_create_material_type(self):
        self.client.force_login(self.admin_user)

        response = self.client.post(
            "/party/admin/rules/",
            {
                "operation": "save_material",
                "track_type": PartyMemberStatus.TrackType.PARTY,
                "current_stage_id": self.stage_one.id,
                "code": "training_certificate",
                "name": "培训结业证书",
                "allowed_extensions": ".pdf, docx ,PDF",
                "max_size_mb": 10,
                "is_active": "on",
                "description": "培训材料",
            },
        )

        self.assertEqual(response.status_code, 302)
        material = PartyMaterialType.objects.get(code="training_certificate")
        self.assertEqual(material.allowed_extensions, "docx,pdf")
        self.assertTrue(
            AuditLog.objects.filter(
                user=self.admin_user,
                action="save party material type",
                target_id=str(material.id),
            ).exists()
        )

    def test_admin_rules_page_can_create_stage_material_requirement(self):
        self.client.force_login(self.admin_user)

        response = self.client.post(
            "/party/admin/rules/",
            {
                "operation": "save_requirement",
                "track_type": PartyMemberStatus.TrackType.PARTY,
                "current_stage_id": self.stage_two.id,
                "stage_definition": self.stage_two.id,
                "material_type": self.material_type.id,
                "is_required": "on",
                "order": 1,
                "note": "进入积极分子阶段前必须补齐",
            },
        )

        self.assertEqual(response.status_code, 302)
        requirement = PartyStageMaterialRequirement.objects.get(
            stage_definition=self.stage_two,
            material_type=self.material_type,
        )
        self.assertTrue(requirement.is_required)
        self.assertEqual(requirement.note, "进入积极分子阶段前必须补齐")
        self.assertTrue(
            AuditLog.objects.filter(
                user=self.admin_user,
                action="save party stage requirement",
                target_id=str(requirement.id),
            ).exists()
        )

    def test_submit_approval_creates_approval_pending_reminder(self):
        self._submit_required_material()

        workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="提交审批",
        )

        approval_reminder = PartyReminder.objects.get(
            member=self.member,
            stage_instance=self.current_stage,
            reminder_type=PartyReminder.ReminderType.APPROVAL_PENDING,
        )
        self.assertEqual(approval_reminder.status, PartyReminder.Status.PENDING)
        self.assertIn("等待老师处理", approval_reminder.content)

    def test_generating_reminders_for_open_stage_creates_material_and_ready_reminders(self):
        self.current_stage.due_at = timezone.now() + timezone.timedelta(days=2)
        self.current_stage.save(update_fields=["due_at"])

        reminder.sync_member_reminders(self.member)

        reminder_types = set(
            PartyReminder.objects.filter(member=self.member).values_list(
                "reminder_type",
                flat=True,
            )
        )
        self.assertIn(PartyReminder.ReminderType.MATERIAL_DUE, reminder_types)
        self.assertIn(PartyReminder.ReminderType.STAGE_READY, reminder_types)

    def test_generate_party_reminders_command_builds_missing_reminders(self):
        self.current_stage.due_at = timezone.now() + timezone.timedelta(days=1)
        self.current_stage.save(update_fields=["due_at"])
        PartyReminder.objects.all().delete()

        call_command("generate_party_reminders", member_id=self.member.id)

        self.assertTrue(
            PartyReminder.objects.filter(
                member=self.member,
                reminder_type=PartyReminder.ReminderType.STAGE_READY,
            ).exists()
        )

    def test_send_party_reminders_command_marks_due_reminders_as_sent(self):
        PartyReminder.objects.create(
            member=self.member,
            stage_instance=self.current_stage,
            reminder_type=PartyReminder.ReminderType.APPROVAL_PENDING,
            channel=PartyReminder.Channel.SITE,
            status=PartyReminder.Status.PENDING,
            title="审批提醒",
            content="请关注审批结果",
            scheduled_at=timezone.now() - timezone.timedelta(minutes=1),
        )

        call_command("send_party_reminders")

        sent_reminder = PartyReminder.objects.get(
            member=self.member,
            reminder_type=PartyReminder.ReminderType.APPROVAL_PENDING,
        )
        self.assertEqual(sent_reminder.status, PartyReminder.Status.SENT)
        self.assertIsNotNone(sent_reminder.sent_at)

    def test_student_dashboard_shows_reminders(self):
        PartyReminder.objects.create(
            member=self.member,
            stage_instance=self.current_stage,
            reminder_type=PartyReminder.ReminderType.APPROVAL_PENDING,
            channel=PartyReminder.Channel.SITE,
            status=PartyReminder.Status.PENDING,
            title="审批处理中",
            content="你的当前节点正在等待老师审批。",
            scheduled_at=timezone.now(),
        )

        self.client.force_login(self.student_user)
        response = self.client.get("/party/student/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "提醒事项")
        self.assertContains(response, "审批处理中")
        self.assertContains(response, "你的当前节点正在等待老师审批")
        self.assertContains(response, "返回党团模块入口")

    def test_student_reminder_page_lists_all_reminders(self):
        PartyReminder.objects.create(
            member=self.member,
            stage_instance=self.current_stage,
            reminder_type=PartyReminder.ReminderType.APPROVAL_PENDING,
            channel=PartyReminder.Channel.SITE,
            status=PartyReminder.Status.PENDING,
            title="审批处理中",
            content="你的当前节点正在等待老师审批。",
            scheduled_at=timezone.now(),
        )
        PartyReminder.objects.create(
            member=self.member,
            stage_instance=self.current_stage,
            reminder_type=PartyReminder.ReminderType.MATERIAL_DUE,
            channel=PartyReminder.Channel.SITE,
            status=PartyReminder.Status.SENT,
            title="材料准备提醒",
            content="请尽快补齐当前节点材料。",
            scheduled_at=timezone.now() - timezone.timedelta(days=1),
            sent_at=timezone.now() - timezone.timedelta(hours=1),
        )

        self.client.force_login(self.student_user)
        response = self.client.get("/party/student/reminders/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "提醒与待办")
        self.assertContains(response, "审批处理中")
        self.assertContains(response, "材料准备提醒")
        self.assertContains(response, "待处理提醒")
        self.assertContains(response, "已发送提醒")

    def test_student_subpages_have_back_actions(self):
        self.client.force_login(self.student_user)

        response = self.client.get("/party/student/timeline/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "返回学生端总览")
        self.assertContains(response, "返回党团模块入口")

    def test_student_timeline_page_shows_breadcrumb_navigation(self):
        self.client.force_login(self.student_user)

        response = self.client.get("/party/student/timeline/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "breadcrumb-list")
        self.assertContains(response, reverse("party:party_student:dashboard"))

    def test_dashboard_subpage_links_preserve_return_to(self):
        self.client.force_login(self.student_user)

        response = self.client.get("/party/student/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "?return_to=/party/student/")

    def test_admin_import_export_page_renders(self):
        self.client.force_login(self.admin_user)

        response = self.client.get("/party/admin/import/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "批量导入导出")
        self.assertContains(response, "下载模板")
        self.assertContains(response, "导出 Excel")
        self.assertContains(response, "返回管理端总览")

    def test_admin_reminder_overview_page_renders_counts_and_overdue_stage(self):
        PartyReminder.objects.create(
            member=self.member,
            stage_instance=self.current_stage,
            reminder_type=PartyReminder.ReminderType.APPROVAL_PENDING,
            channel=PartyReminder.Channel.SITE,
            status=PartyReminder.Status.PENDING,
            title="审批提醒",
            content="请关注审批结果",
            scheduled_at=timezone.now(),
        )
        self.current_stage.due_at = timezone.now() - timezone.timedelta(days=1)
        self.current_stage.save(update_fields=["due_at"])

        self.client.force_login(self.admin_user)
        response = self.client.get("/party/admin/reminders/")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "提醒概览")
        self.assertContains(response, "待发送提醒")
        self.assertContains(response, "逾期节点")
        self.assertContains(response, "审批提醒")
        self.assertContains(response, "已逾期")

    def test_admin_import_export_page_downloads_template(self):
        self.client.force_login(self.admin_user)

        response = self.client.get("/party/admin/import/?action=download_template")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content))
        worksheet = workbook.active
        headers = [cell.value for cell in worksheet[1]]
        self.assertIn("username", headers)
        self.assertIn("completed_stage_definition_codes", headers)
        self.assertIn("current_stage_definition_code", headers)

    def test_admin_import_export_page_exports_member_status(self):
        self.client.force_login(self.admin_user)

        response = self.client.get("/party/admin/import/?action=export_members")

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        worksheet = workbook.active
        exported_rows = list(worksheet.iter_rows(min_row=2, values_only=True))
        self.assertTrue(any(row[0] == "student_user" for row in exported_rows))

    def test_admin_import_export_page_exports_approval_details(self):
        self._submit_required_material()
        task = workflow.submit_stage_for_approval(
            self.current_stage,
            self.student_user,
            comment="待审批",
        )
        approval.reject_task(task, self.admin_user, comment="请补材料")

        self.client.force_login(self.admin_user)
        response = self.client.get("/party/admin/import/?action=export_approvals")

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        self.assertIn("approval_tasks", workbook.sheetnames)
        self.assertIn("approval_actions", workbook.sheetnames)
        task_rows = list(workbook["approval_tasks"].iter_rows(min_row=2, values_only=True))
        action_rows = list(workbook["approval_actions"].iter_rows(min_row=2, values_only=True))
        self.assertTrue(any(row[0] == task.id for row in task_rows))
        self.assertTrue(any(row[0] == task.id and row[1] == "reject" for row in action_rows))

    def test_admin_import_export_page_imports_member_workbook(self):
        self.client.force_login(self.admin_user)
        workbook = load_workbook(BytesIO(self.client.get("/party/admin/import/?action=download_template").content))
        worksheet = workbook.active
        worksheet.delete_rows(2, worksheet.max_row)
        worksheet.append(
            [
                "import_student",
                "20269999",
                "导入学生",
                "2026",
                "数据科学",
                "party",
                "activist",
                "2026-04-01",
                "2026-07-01",
                False,
                "批量导入测试",
                "party_applicant",
                "party_activist",
                "open",
                "2026-04-01 09:00:00",
                "2026-07-01 09:00:00",
            ]
        )
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        response = self.client.post(
            "/party/admin/import/",
            {
                "operation": "preview_import",
                "file": SimpleUploadedFile(
                    "party_import.xlsx",
                    buffer.getvalue(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

        self.assertEqual(response.status_code, 302)
        preview_response = self.client.get("/party/admin/import/")
        self.assertEqual(preview_response.status_code, 200)
        self.assertContains(preview_response, "导入预校验结果")
        self.assertContains(preview_response, "确认导入写库")

        response = self.client.post(
            "/party/admin/import/",
            {
                "operation": "confirm_import",
            },
        )

        self.assertEqual(response.status_code, 302)
        imported_user = User.objects.get(username="import_student")
        imported_member = PartyMemberStatus.objects.get(user=imported_user)
        self.assertEqual(imported_member.current_stage, PartyMemberStatus.Stage.ACTIVIST)
        imported_stage = PartyStageInstance.objects.get(
            member=imported_member,
            stage_definition__code="party_activist",
        )
        self.assertEqual(imported_stage.status, PartyStageInstance.Status.OPEN)
        historical_stage = PartyStageInstance.objects.get(
            member=imported_member,
            stage_definition__code="party_applicant",
        )
        self.assertEqual(historical_stage.status, PartyStageInstance.Status.APPROVED)
        self.assertTrue(
            AuditLog.objects.filter(
                user=self.admin_user,
                action="import party member workbook",
            ).exists()
        )

    def test_admin_import_export_page_provides_error_report_download(self):
        self.client.force_login(self.admin_user)
        PartyStageDefinition.objects.create(
            code="party_candidate",
            name="发展对象",
            track_type=PartyMemberStatus.TrackType.PARTY,
            stage=PartyMemberStatus.Stage.CANDIDATE,
            order=3,
            target_role_for_approval=User.ROLE_ADMIN,
        )
        workbook = load_workbook(
            BytesIO(self.client.get("/party/admin/import/?action=download_template").content)
        )
        worksheet = workbook.active
        worksheet.delete_rows(2, worksheet.max_row)
        worksheet.append(
            [
                "broken_student",
                "20268888",
                "错误学生",
                "2026",
                "人工智能",
                "party",
                "activist",
                "2026-04-01",
                "2026-07-01",
                False,
                "错误导入测试",
                "party_candidate",
                "party_activist",
                "open",
                "2026-04-01 09:00:00",
                "2026-07-01 09:00:00",
            ]
        )
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        response = self.client.post(
            "/party/admin/import/",
            {
                "operation": "preview_import",
                "file": SimpleUploadedFile(
                    "party_import_invalid.xlsx",
                    buffer.getvalue(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "下载错误清单")

        error_response = self.client.get("/party/admin/import/?action=download_errors")
        self.assertEqual(error_response.status_code, 200)
        workbook = load_workbook(BytesIO(error_response.content))
        worksheet = workbook.active
        error_rows = list(worksheet.iter_rows(min_row=2, values_only=True))
        self.assertTrue(any("历史已完成节点必须早于当前节点" in row[1] for row in error_rows))

    def test_admin_import_export_page_preview_can_be_cleared(self):
        self.client.force_login(self.admin_user)
        workbook = load_workbook(
            BytesIO(self.client.get("/party/admin/import/?action=download_template").content)
        )
        worksheet = workbook.active
        worksheet.delete_rows(2, worksheet.max_row)
        worksheet.append(
            [
                "preview_student",
                "20267777",
                "预览学生",
                "2026",
                "软件工程",
                "party",
                "applicant",
                "2026-04-01",
                "2026-07-01",
                False,
                "预览测试",
                "",
                "party_applicant",
                "open",
                "2026-04-01 09:00:00",
                "2026-07-01 09:00:00",
            ]
        )
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        response = self.client.post(
            "/party/admin/import/",
            {
                "operation": "preview_import",
                "file": SimpleUploadedFile(
                    "party_import_preview.xlsx",
                    buffer.getvalue(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        self.assertEqual(response.status_code, 302)

        response = self.client.get("/party/admin/import/?action=clear_preview")
        self.assertEqual(response.status_code, 302)
        final_page = self.client.get("/party/admin/import/")
        self.assertNotContains(final_page, "确认导入写库")


class PartyLeagueWorkflowTestCase(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._temp_media = Path(settings.BASE_DIR) / ".tmp_party_test_media_league"
        cls._temp_media.mkdir(parents=True, exist_ok=True)
        cls._override = override_settings(MEDIA_ROOT=cls._temp_media)
        cls._override.enable()

    @classmethod
    def tearDownClass(cls):
        cls._override.disable()
        shutil.rmtree(cls._temp_media, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.admin_user = User.objects.create_user(
            username="league_admin",
            password="party1234",
            role=User.ROLE_ADMIN,
            real_name="团务老师",
        )
        self.student_user = User.objects.create_user(
            username="league_student",
            password="party1234",
            role=User.ROLE_STUDENT,
            real_name="团员学生",
            student_id="20260101",
        )
        self.material_type = PartyMaterialType.objects.create(
            code="league_application",
            name="入团申请书",
            allowed_extensions="pdf,doc,docx",
            max_size_mb=30,
        )
        self.stage_one = PartyStageDefinition.objects.create(
            code="league_applicant",
            name="入团申请人",
            track_type=PartyMemberStatus.TrackType.LEAGUE,
            stage=PartyMemberStatus.Stage.APPLICANT,
            order=1,
            target_role_for_approval=User.ROLE_ADMIN,
        )
        self.stage_two = PartyStageDefinition.objects.create(
            code="league_activist",
            name="入团积极分子",
            track_type=PartyMemberStatus.TrackType.LEAGUE,
            stage=PartyMemberStatus.Stage.ACTIVIST,
            order=2,
            min_duration_days=30,
            target_role_for_approval=User.ROLE_ADMIN,
        )
        PartyStageMaterialRequirement.objects.create(
            stage_definition=self.stage_one,
            material_type=self.material_type,
            is_required=True,
            order=1,
        )
        self.member = PartyMemberStatus.objects.create(
            user=self.student_user,
            track_type=PartyMemberStatus.TrackType.LEAGUE,
            current_stage=PartyMemberStatus.Stage.APPLICANT,
            created_by=self.admin_user,
        )
        self.current_stage = workflow.initialize_member_workflow(
            self.member,
            operator=self.admin_user,
        )

    def _submit_required_material(self):
        upload = SimpleUploadedFile(
            "league_application.pdf",
            b"league content",
            content_type="application/pdf",
        )
        return PartyMaterialSubmission.objects.create(
            stage_instance=self.current_stage,
            material_type=self.material_type,
            file=upload,
            version=workflow.get_next_material_version(self.current_stage, self.material_type),
            submitted_by=self.student_user,
        )

    def test_league_initialize_member_workflow_uses_league_first_stage(self):
        self.assertEqual(self.current_stage.stage_definition.track_type, PartyMemberStatus.TrackType.LEAGUE)
        self.assertEqual(self.current_stage.stage_definition.code, "league_applicant")
        self.assertEqual(self.member.current_stage, PartyMemberStatus.Stage.APPLICANT)

    def test_league_approval_advances_to_next_league_stage(self):
        self._submit_required_material()
        task = workflow.submit_stage_for_approval(self.current_stage, self.student_user)

        next_stage = approval.approve_task(task, self.admin_user, comment="入团流程通过")

        task.refresh_from_db()
        self.current_stage.refresh_from_db()
        self.member.refresh_from_db()
        self.assertEqual(task.status, PartyApprovalTask.Status.APPROVED)
        self.assertEqual(self.current_stage.status, PartyStageInstance.Status.APPROVED)
        self.assertIsNotNone(next_stage)
        self.assertEqual(next_stage.stage_definition.code, "league_activist")
        self.assertEqual(next_stage.stage_definition.track_type, PartyMemberStatus.TrackType.LEAGUE)
        self.assertEqual(next_stage.status, PartyStageInstance.Status.OPEN)
        self.assertEqual(self.member.current_stage, PartyMemberStatus.Stage.ACTIVIST)

    def test_seed_party_data_initializes_league_stage_definitions(self):
        self.member.stage_instances.all().delete()
        self.member.delete()
        PartyStageMaterialRequirement.objects.all().delete()
        PartyStageDefinition.objects.all().delete()
        PartyMaterialType.objects.all().delete()

        call_command("seed_party_data")

        self.assertTrue(
            PartyStageDefinition.objects.filter(
                track_type=PartyMemberStatus.TrackType.LEAGUE,
                code="league_applicant",
            ).exists()
        )
        self.assertTrue(
            PartyStageDefinition.objects.filter(
                track_type=PartyMemberStatus.TrackType.LEAGUE,
                code="league_full",
            ).exists()
        )
        self.assertTrue(
            PartyStageMaterialRequirement.objects.filter(
                stage_definition__track_type=PartyMemberStatus.TrackType.LEAGUE,
                material_type__code="league_application",
            ).exists()
        )

    def test_seed_party_data_with_demo_users_creates_league_demo_member(self):
        self.member.stage_instances.all().delete()
        self.member.delete()
        self.student_user.delete()
        self.admin_user.delete()
        PartyStageMaterialRequirement.objects.all().delete()
        PartyStageDefinition.objects.all().delete()
        PartyMaterialType.objects.all().delete()
        User.objects.filter(username="league_student").delete()
        User.objects.filter(username="party_student").delete()
        User.objects.filter(username="party_admin").delete()

        call_command("seed_party_data", with_demo_users=True)

        demo_admin = User.objects.get(username="party_admin")
        league_student = User.objects.get(username="league_student")
        party_student = User.objects.get(username="party_student")
        league_member = PartyMemberStatus.objects.get(user=league_student)
        party_member = PartyMemberStatus.objects.get(user=party_student)

        self.assertEqual(demo_admin.role, User.ROLE_ADMIN)
        self.assertEqual(league_member.track_type, PartyMemberStatus.TrackType.LEAGUE)
        self.assertEqual(party_member.track_type, PartyMemberStatus.TrackType.PARTY)
        self.assertTrue(
            PartyStageInstance.objects.filter(
                member=league_member,
                stage_definition__track_type=PartyMemberStatus.TrackType.LEAGUE,
                status=PartyStageInstance.Status.OPEN,
            ).exists()
        )
