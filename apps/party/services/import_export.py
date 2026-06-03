from datetime import date, datetime, time, timedelta
from io import BytesIO

from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime
from openpyxl import Workbook, load_workbook

from apps.users.models import User

from ..models import (
    PartyApprovalTask,
    PartyMemberStatus,
    PartyReminder,
    PartyStageDefinition,
    PartyStageInstance,
)
from . import reminder, workflow


class PartyImportValidationError(Exception):
    def __init__(self, error_rows):
        super().__init__("导入数据存在错误。")
        self.error_rows = error_rows


IMPORT_HEADERS = [
    "username",
    "student_id",
    "real_name",
    "grade",
    "major",
    "track_type",
    "current_stage",
    "status_since",
    "expected_next_date",
    "is_locked",
    "notes",
    "completed_stage_definition_codes",
    "current_stage_definition_code",
    "current_stage_status",
    "current_stage_started_at",
    "current_stage_due_at",
]


EXPORT_HEADERS = IMPORT_HEADERS + [
    "pending_task_count",
    "pending_reminder_count",
]


APPROVAL_TASK_HEADERS = [
    "task_id",
    "member_id",
    "username",
    "real_name",
    "stage_code",
    "stage_name",
    "task_status",
    "assignee_role",
    "assignee",
    "created_at",
    "processed_at",
    "can_withdraw_until",
    "comment",
]


APPROVAL_ACTION_HEADERS = [
    "task_id",
    "action",
    "operator",
    "created_at",
    "comment",
]


def _serialize_cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _serialize_error_rows(error_rows):
    serialized = []
    for item in error_rows:
        serialized.append(
            {
                "row_number": item["row_number"],
                "message": item["message"],
                "row_data": {
                    key: _serialize_cell(value) for key, value in item["row_data"].items()
                },
            }
        )
    return serialized


def _serialize_row_data(row_data):
    return {key: _serialize_cell(value) for key, value in row_data.items()}


def _normalize_bool(value):
    if isinstance(value, bool):
        return value
    if value in (None, ""):
        return False
    text = str(value).strip().lower()
    return text in {"1", "true", "yes", "y", "是"}


def _normalize_date(value, field_name):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        parsed = value.date()
        if parsed.year <= 1:
            return None
        return parsed
    if isinstance(value, date):
        if value.year <= 1:
            return None
        return value
    parsed = parse_date(str(value).strip())
    if parsed is None:
        raise ValidationError(f"{field_name} 日期格式无效，需为 YYYY-MM-DD。")
    if parsed.year <= 1:
        return None
    return parsed


def _normalize_datetime(value, field_name):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        dt = value
    else:
        dt = parse_datetime(str(value).strip())
        if dt is None:
            raise ValidationError(f"{field_name} 时间格式无效，需为 YYYY-MM-DD HH:MM[:SS]。")
    if dt.year <= 1:
        return None
    if timezone.is_naive(dt):
        dt = timezone.make_aware(dt, timezone.get_current_timezone())
    return dt


def _parse_code_list(value):
    if value in (None, ""):
        return []
    seen = set()
    codes = []
    for item in str(value).split(","):
        code = item.strip()
        if code and code not in seen:
            seen.add(code)
            codes.append(code)
    return codes


def _get_stage_definition(track_type, code):
    if not code:
        return None
    try:
        return PartyStageDefinition.objects.get(track_type=track_type, code=code)
    except PartyStageDefinition.DoesNotExist as exc:
        raise ValidationError(f"未找到节点定义：{track_type}/{code}。") from exc


def _get_stage_definitions(track_type, codes):
    definitions = []
    for code in codes:
        definitions.append(_get_stage_definition(track_type, code))
    return definitions


def _get_or_create_student(row):
    username = str(row.get("username") or "").strip()
    student_id = str(row.get("student_id") or "").strip()
    real_name = str(row.get("real_name") or "").strip()

    if not username:
        raise ValidationError("username 不能为空。")

    user = User.objects.filter(username=username).first()
    if user is None and student_id:
        user = User.objects.filter(student_id=student_id).first()
    if user is None:
        user = User(username=username)
        user.set_unusable_password()

    user.role = User.ROLE_STUDENT
    user.student_id = student_id or user.student_id
    user.real_name = real_name or user.real_name
    user.grade = str(row.get("grade") or "").strip()
    user.major = str(row.get("major") or "").strip()
    user.save()
    return user


def _collect_existing_student(row):
    username = str(row.get("username") or "").strip()
    student_id = str(row.get("student_id") or "").strip()
    user = User.objects.filter(username=username).first()
    if user is None and student_id:
        user = User.objects.filter(student_id=student_id).first()
    return user


def _build_reference_start(member, row):
    current_started_at = _normalize_datetime(
        row.get("current_stage_started_at"),
        "current_stage_started_at",
    )
    if current_started_at is not None:
        return current_started_at

    status_since = _normalize_date(row.get("status_since"), "status_since")
    if status_since is not None:
        return timezone.make_aware(
            datetime.combine(status_since, time(hour=9)),
            timezone.get_current_timezone(),
        )

    return timezone.now()


def _sync_historical_completed_stages(member, row, current_stage_definition):
    completed_codes = _parse_code_list(row.get("completed_stage_definition_codes"))
    if not completed_codes:
        return []

    completed_definitions = _get_stage_definitions(member.track_type, completed_codes)
    if current_stage_definition is not None:
        invalid_codes = [
            definition.code
            for definition in completed_definitions
            if definition.order >= current_stage_definition.order
        ]
        if invalid_codes:
            raise ValidationError(
                "历史已完成节点必须早于当前节点，以下编码不合法："
                + ", ".join(invalid_codes)
            )

    reference_started_at = _build_reference_start(member, row)
    created_instances = []
    total = len(completed_definitions)
    for index, definition in enumerate(sorted(completed_definitions, key=lambda item: item.order)):
        delta_days = total - index
        completed_at = reference_started_at - timedelta(days=delta_days)
        started_at = completed_at - timedelta(days=max(definition.min_duration_days, 0))
        stage_instance, _ = PartyStageInstance.objects.get_or_create(
            member=member,
            stage_definition=definition,
        )
        stage_instance.status = PartyStageInstance.Status.APPROVED
        stage_instance.started_at = started_at
        stage_instance.due_at = completed_at if definition.min_duration_days else None
        stage_instance.completed_at = completed_at
        stage_instance.remark = "由 Excel 批量导入初始化为历史已完成节点"
        stage_instance.save()
        created_instances.append(stage_instance)
    return created_instances


def _sync_current_stage_instance(member, row):
    stage_code = str(row.get("current_stage_definition_code") or "").strip()
    if not stage_code:
        return None

    stage_definition = _get_stage_definition(member.track_type, stage_code)
    status = str(row.get("current_stage_status") or PartyStageInstance.Status.OPEN).strip()
    valid_statuses = {value for value, _ in PartyStageInstance.Status.choices}
    if status not in valid_statuses:
        raise ValidationError(f"current_stage_status 无效：{status}")

    started_at = _normalize_datetime(row.get("current_stage_started_at"), "current_stage_started_at")
    due_at = _normalize_datetime(row.get("current_stage_due_at"), "current_stage_due_at")
    completed_at = None
    if status == PartyStageInstance.Status.APPROVED:
        completed_at = started_at or timezone.now()

    stage_instance, _ = PartyStageInstance.objects.get_or_create(
        member=member,
        stage_definition=stage_definition,
    )
    stage_instance.status = status
    stage_instance.started_at = started_at
    stage_instance.due_at = due_at
    stage_instance.completed_at = completed_at
    stage_instance.save()
    return stage_instance


def _parse_workbook_rows(uploaded_file):
    workbook = load_workbook(uploaded_file, data_only=True)
    worksheet = workbook.active
    headers = [cell.value for cell in worksheet[1]]
    if headers != IMPORT_HEADERS:
        raise ValidationError("导入模板表头不匹配，请先下载模板后再填写。")

    rows = []
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        if not any(value not in (None, "") for value in row):
            continue
        row_data = dict(zip(IMPORT_HEADERS, row))
        rows.append({"row_number": row_number, "row_data": row_data})
    return rows


def _validate_row_for_preview(row_data):
    username = str(row_data.get("username") or "").strip()
    if not username:
        raise ValidationError("username 不能为空。")

    track_type = str(row_data.get("track_type") or PartyMemberStatus.TrackType.PARTY).strip()
    valid_track_types = {value for value, _ in PartyMemberStatus.TrackType.choices}
    if track_type not in valid_track_types:
        raise ValidationError(f"track_type 无效：{track_type}")

    current_stage = str(
        row_data.get("current_stage") or PartyMemberStatus.Stage.APPLICANT
    ).strip()
    valid_stages = {value for value, _ in PartyMemberStatus.Stage.choices}
    if current_stage not in valid_stages:
        raise ValidationError(f"current_stage 无效：{current_stage}")

    _normalize_date(row_data.get("status_since"), "status_since")
    _normalize_date(row_data.get("expected_next_date"), "expected_next_date")
    _normalize_bool(row_data.get("is_locked"))

    current_stage_definition_code = str(row_data.get("current_stage_definition_code") or "").strip()
    current_stage_definition = None
    if current_stage_definition_code:
        current_stage_definition = _get_stage_definition(track_type, current_stage_definition_code)

    current_stage_status = str(
        row_data.get("current_stage_status") or PartyStageInstance.Status.OPEN
    ).strip()
    valid_statuses = {value for value, _ in PartyStageInstance.Status.choices}
    if current_stage_status not in valid_statuses:
        raise ValidationError(f"current_stage_status 无效：{current_stage_status}")

    _normalize_datetime(row_data.get("current_stage_started_at"), "current_stage_started_at")
    _normalize_datetime(row_data.get("current_stage_due_at"), "current_stage_due_at")

    completed_codes = _parse_code_list(row_data.get("completed_stage_definition_codes"))
    completed_definitions = _get_stage_definitions(track_type, completed_codes)
    if current_stage_definition is not None:
        invalid_codes = [
            definition.code
            for definition in completed_definitions
            if definition.order >= current_stage_definition.order
        ]
        if invalid_codes:
            raise ValidationError(
                "历史已完成节点必须早于当前节点，以下编码不合法：" + ", ".join(invalid_codes)
            )

    existing_user = _collect_existing_student(row_data)
    existing_member = None
    if existing_user is not None:
        existing_member = PartyMemberStatus.objects.filter(user=existing_user).first()

    return {
        "username": username,
        "student_id": str(row_data.get("student_id") or "").strip(),
        "real_name": str(row_data.get("real_name") or "").strip(),
        "track_type": track_type,
        "current_stage": current_stage,
        "completed_stage_codes": [definition.code for definition in completed_definitions],
        "current_stage_definition_code": current_stage_definition.code if current_stage_definition else "",
        "current_stage_status": current_stage_status,
        "student_action": "更新" if existing_user else "新增",
        "member_action": "更新" if existing_member else "新增",
    }


def preview_member_workbook(uploaded_file):
    parsed_rows = _parse_workbook_rows(uploaded_file)
    preview_rows = []
    error_rows = []
    for item in parsed_rows:
        row_data = item["row_data"]
        try:
            preview = _validate_row_for_preview(row_data)
            preview_rows.append(
                {
                    "row_number": item["row_number"],
                    "row_data": _serialize_row_data(row_data),
                    "preview": preview,
                }
            )
        except ValidationError as exc:
            error_rows.append(
                {
                    "row_number": item["row_number"],
                    "message": "; ".join(exc.messages) if hasattr(exc, "messages") else str(exc),
                    "row_data": row_data,
                }
            )

    if error_rows:
        raise PartyImportValidationError(_serialize_error_rows(error_rows))
    return preview_rows


def _import_row_data(row_data, operator):
    user = _get_or_create_student(row_data)
    track_type = str(row_data.get("track_type") or PartyMemberStatus.TrackType.PARTY).strip()
    current_stage = str(
        row_data.get("current_stage") or PartyMemberStatus.Stage.APPLICANT
    ).strip()

    member, _ = PartyMemberStatus.objects.get_or_create(
        user=user,
        defaults={
            "track_type": track_type,
            "current_stage": current_stage,
            "created_by": operator,
        },
    )
    member.track_type = track_type
    member.current_stage = current_stage
    member.status_since = _normalize_date(row_data.get("status_since"), "status_since")
    member.expected_next_date = _normalize_date(
        row_data.get("expected_next_date"),
        "expected_next_date",
    )
    member.is_locked = _normalize_bool(row_data.get("is_locked"))
    member.notes = str(row_data.get("notes") or "").strip()
    if member.created_by_id is None:
        member.created_by = operator
    member.save()

    current_stage_definition_code = str(
        row_data.get("current_stage_definition_code") or ""
    ).strip()
    current_stage_definition = _get_stage_definition(
        member.track_type,
        current_stage_definition_code,
    ) if current_stage_definition_code else None

    _sync_historical_completed_stages(member, row_data, current_stage_definition)
    stage_instance = _sync_current_stage_instance(member, row_data)
    if stage_instance is None:
        workflow.initialize_member_workflow(member, operator=operator)
    reminder.sync_member_reminders(member)
    return member


@transaction.atomic
def import_member_workbook(uploaded_file, operator):
    preview_rows = preview_member_workbook(uploaded_file)
    return import_member_preview_rows(preview_rows, operator)


@transaction.atomic
def import_member_preview_rows(preview_rows, operator):
    imported_count = 0
    for item in preview_rows:
        _import_row_data(item["row_data"], operator)
        imported_count += 1
    return imported_count


def build_import_template_response():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "party_import_template"
    worksheet.append(IMPORT_HEADERS)
    worksheet.append(
        [
            "party_student_02",
            "20260002",
            "示例学生二",
            "2026",
            "信息安全",
            "party",
            "activist",
            "2026-04-20",
            "2026-07-20",
            False,
            "示例备注",
            "party_applicant",
            "party_activist",
            "open",
            "2026-04-20 09:00:00",
            "2026-07-20 09:00:00",
        ]
    )
    worksheet.append(
        [
            "league_student_02",
            "20260003",
            "示例团员学生",
            "2026",
            "网络空间安全",
            "league",
            "activist",
            "2026-04-10",
            "2026-05-10",
            False,
            "示例入团流程导入",
            "league_applicant",
            "league_activist",
            "open",
            "2026-04-10 09:00:00",
            "2026-05-10 09:00:00",
        ]
    )
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="party_import_template.xlsx"'
    return response


def build_import_error_response(error_rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "party_import_errors"
    worksheet.append(["row_number", "message"] + IMPORT_HEADERS)
    for item in error_rows:
        worksheet.append(
            [item.get("row_number"), item.get("message")]
            + [item.get("row_data", {}).get(header, "") for header in IMPORT_HEADERS]
        )
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="party_import_errors.xlsx"'
    return response


def build_member_export_response():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "party_members"
    worksheet.append(EXPORT_HEADERS)

    members = PartyMemberStatus.objects.select_related("user").prefetch_related(
        "stage_instances__stage_definition",
        "reminders",
    )
    for member in members:
        current_stage_instance = workflow.get_current_stage_instance(member)
        completed_codes = ",".join(
            member.stage_instances.filter(status=PartyStageInstance.Status.APPROVED)
            .select_related("stage_definition")
            .order_by("stage_definition__order")
            .values_list("stage_definition__code", flat=True)
        )
        worksheet.append(
            [
                member.user.username,
                member.user.student_id,
                member.user.real_name,
                member.user.grade,
                member.user.major,
                member.track_type,
                member.current_stage,
                member.status_since.isoformat() if member.status_since else "",
                member.expected_next_date.isoformat() if member.expected_next_date else "",
                member.is_locked,
                member.notes,
                completed_codes,
                current_stage_instance.stage_definition.code if current_stage_instance else "",
                current_stage_instance.status if current_stage_instance else "",
                (
                    timezone.localtime(current_stage_instance.started_at).strftime("%Y-%m-%d %H:%M:%S")
                    if current_stage_instance and current_stage_instance.started_at
                    else ""
                ),
                (
                    timezone.localtime(current_stage_instance.due_at).strftime("%Y-%m-%d %H:%M:%S")
                    if current_stage_instance and current_stage_instance.due_at
                    else ""
                ),
                PartyApprovalTask.objects.filter(
                    stage_instance__member=member,
                    status=PartyApprovalTask.Status.PENDING,
                ).count(),
                PartyReminder.objects.filter(
                    member=member,
                    status=PartyReminder.Status.PENDING,
                ).count(),
            ]
        )

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="party_member_status_export.xlsx"'
    return response


def build_approval_export_response():
    workbook = Workbook()
    task_sheet = workbook.active
    task_sheet.title = "approval_tasks"
    task_sheet.append(APPROVAL_TASK_HEADERS)

    action_sheet = workbook.create_sheet("approval_actions")
    action_sheet.append(APPROVAL_ACTION_HEADERS)

    tasks = PartyApprovalTask.objects.select_related(
        "stage_instance__member__user",
        "stage_instance__stage_definition",
        "assignee",
    ).prefetch_related("actions__operator")

    for task in tasks:
        task_sheet.append(
            [
                task.id,
                task.stage_instance.member_id,
                task.stage_instance.member.user.username,
                task.stage_instance.member.user.real_name,
                task.stage_instance.stage_definition.code,
                task.stage_instance.stage_definition.name,
                task.status,
                task.get_assignee_role_display(),
                task.assignee.real_name if task.assignee else "",
                timezone.localtime(task.created_at).strftime("%Y-%m-%d %H:%M:%S"),
                (
                    timezone.localtime(task.processed_at).strftime("%Y-%m-%d %H:%M:%S")
                    if task.processed_at
                    else ""
                ),
                (
                    timezone.localtime(task.can_withdraw_until).strftime("%Y-%m-%d %H:%M:%S")
                    if task.can_withdraw_until
                    else ""
                ),
                task.comment,
            ]
        )
        for action in task.actions.all().order_by("created_at"):
            action_sheet.append(
                [
                    task.id,
                    action.action,
                    action.operator.real_name if action.operator else "",
                    timezone.localtime(action.created_at).strftime("%Y-%m-%d %H:%M:%S"),
                    action.comment,
                ]
            )

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="party_approval_detail_export.xlsx"'
    return response
